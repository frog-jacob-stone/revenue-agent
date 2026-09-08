"""Template in, .docx bytes out. Pure — touches no database.

The whole reason this feature renders an existing document rather than building
one is formatting fidelity: the output has to be indistinguishable from the
template a human maintains in Word. Three decisions follow from that, and all
three were arrived at by measuring rather than by reading docs. They are easy to
undo by accident, so each says what breaks.

**Plain `{{ tag }}`, never `{{r tag }}`.** docxtpl's `RichText` builds its own
`<w:r>` and replaces the tag's entire run with it, discarding that run's
properties. A tag typed inside a bold Times-14pt clause renders unbolded in the
body font — measurably: `bold=True font='Times New Roman' size=14pt` in, `bold=None
font=None size=None` out. A plain tag substitutes text *within* the existing run
and keeps every property. So `RichText` is not used here at all.

**`autoescape=True` on render.** Without it, values are pasted into the XML
unescaped and an `&` or `<` corrupts the file. `"Smith & Co <Holdings>"` rendered
as `"Smith  Co "` and swallowed the following table — the document's structure,
not just the string. `Smith & Co` is an ordinary legal entity name, so this is a
live bug rather than a hypothetical.

**Highlighting happens after rendering, not during it.** Since values are plain
strings they cannot carry formatting, so `_highlight_review_markers` walks the
rendered document and marks the `[REVIEW: …]` tokens. Doing it by hand rather
than via `RichText(highlight=…)` is not incidental: it emits Word's real
`<w:highlight>` (the highlighter pen) instead of `<w:shd>` run shading, and
because the split deep-copies the source run's `rPr`, the marker inherits the
formatting of the text it replaced. A marker inside an italic Georgia clause
stays italic Georgia.
"""
from __future__ import annotations

import copy
import re
from io import BytesIO
from pathlib import Path

from docx import Document
from docx.enum.text import WD_COLOR_INDEX
from docx.oxml.ns import qn
from docx.text.run import Run
from docxtpl import DocxTemplate

from app.services.contracts.fields import TM_FIELDS, review_marker

TEMPLATES_DIR = Path(__file__).parent / "templates"
TM_TEMPLATE = TEMPLATES_DIR / "tm_agreement.docx"

# Matches what `review_marker()` produces. `[^\]]*` rather than `.*?` so a
# marker can never span a `]` and swallow real contract text after it.
_REVIEW_RE = re.compile(r"\[REVIEW: [^\]]*\]")


class TemplateMissing(RuntimeError):
    """The .docx is not on disk.

    Its own type because the remedy is specific and unguessable from an
    IOError: the template is a versioned repo asset, so a missing one means an
    incomplete checkout or an image built without it — not anything the
    operator did.
    """


class TemplateTagInvalid(RuntimeError):
    """A tag in the .docx is not a bare field name.

    Contract templates support one construct — `{{ field_name }}` — and nothing
    else: no filters, no conditionals, no expressions. That is a deliberate
    limit, because the person maintaining these documents is editing a contract
    in Word, not writing Jinja, and a template that can branch is a template
    whose output nobody can predict from reading it.

    Raised with the offending tag and the fix, because the underlying
    `TemplateSyntaxError` says only "expected token 'end of print statement'"
    and names no tag, no field, and no file.
    """


# `{{r foo }}` is docxtpl's rich-text form and is what its own documentation
# steers an author toward, so templates arrive containing it. Accepted and
# normalized away rather than rejected — but *only* by rewriting it to a plain
# tag, never by honouring it: `RichText` would discard the formatting of the run
# the tag sits in, which is the one thing this whole module exists to preserve.
#
# The leading group swallows any XML tags Word left between `{{` and the `r`,
# because this runs before docxtpl's own run-reassembly (see `patch_xml` below)
# and has to survive a paragraph Word split mid-tag. Those tags are captured
# and put back rather than dropped.
_RICH_TEXT_PREFIX = re.compile(r"(\{\{(?:<[^>]*>)*)\s*r\s+")

# What a tag may contain once normalized: a bare Python identifier.
_TAG_BODY = re.compile(r"\{\{(.*?)\}\}", re.DOTALL)


class _ContractTemplate(DocxTemplate):
    """A `DocxTemplate` that tolerates `{{r … }}` and rejects real logic.

    The `r` prefix has to be stripped *before* `super().patch_xml`, and that
    ordering is the whole subtlety here. docxtpl treats a rich-text tag as a
    placeholder for injected XML, so its `for y in [..., "r"]` pass deletes the
    `<w:r>` and `<w:t>` around the tag to make room. Normalizing afterwards
    leaves a plain tag sitting in that hole, and the rendered value lands as raw
    text *between* two empty runs:

        <w:r><w:t xml:space="preserve"/></w:r>[REVIEW: SOW date]<w:r>…

    which Word silently ignores — the value is in the file and invisible on the
    page, the worst of both. Stripping the prefix first means docxtpl never
    recognises a rich-text tag and substitutes inside the `<w:t>` as normal.
    """

    def patch_xml(self, src_xml: str) -> str:
        patched = super().patch_xml(_RICH_TEXT_PREFIX.sub(r"\1 ", src_xml))
        _assert_bare_tags(patched)
        return patched


def _assert_bare_tags(xml: str) -> None:
    for match in _TAG_BODY.finditer(xml):
        body = match.group(1).strip()
        if not body.isidentifier():
            raise TemplateTagInvalid(
                f"Template tag {{{{ {body} }}}} is not a usable field name. "
                f"A contract tag must be a single word with underscores instead "
                f"of spaces — so `{{{{ {body.replace(' ', '_')} }}}}` rather "
                f"than `{{{{ {body} }}}}`. Fix it in Word, then add a matching "
                f"ContractField in app/services/contracts/fields.py. See "
                f"app/services/contracts/templates/README.md."
            )


def template_tags(path: Path | None = None) -> set[str]:
    """Every field tag the template declares.

    Shared by the renderer and `tests/test_contract_template_tags.py` so both
    read the document through the same normalization — a test that saw raw tags
    while the renderer saw normalized ones would pass on a template that does
    not render.
    """
    return set(
        _ContractTemplate(path or TM_TEMPLATE).get_undeclared_template_variables()
    )


def _blank(value: str | None) -> bool:
    """Whitespace-only counts as unfilled.

    A form submits `""` for an untouched input and `"  "` for one someone
    tabbed through. Both mean the same thing, and treating the second as a real
    value would put an invisible blank into a contract instead of a marker.
    """
    return value is None or not value.strip()


def _iter_paragraph_containers(container):
    """Every paragraph-bearing scope inside `container`, tables included.

    Recursive, because a cell may hold a table. Body text, tables, and nested
    tables are covered; text inside shapes and text boxes lives in a different
    part of the XML and is not reached — templates should not put tags there.
    """
    yield container
    for table in container.tables:
        for row in table.rows:
            for cell in row.cells:
                yield from _iter_paragraph_containers(cell)


def _split_run_on_markers(run: Run) -> int:
    """Give each `[REVIEW: …]` token in `run` its own highlighted run.

    A tag is usually typed as its own run, but nothing guarantees it — a
    template may well read `State of {{ client_state }}, USA` in one run.
    Highlighting the whole run there would paint "State of" and ", USA" yellow
    too, which reads as though the boilerplate needs review.

    So the run is rebuilt as a sequence of runs, each a deep copy of the
    original — that copy is what carries `rPr` across, and it is why the marker
    keeps the font, size, and emphasis of the text it replaced. Returns how many
    markers were highlighted.
    """
    text = run.text
    markers = _REVIEW_RE.findall(text)
    if not markers:
        return 0

    between = _REVIEW_RE.split(text)
    # Interleave back into (chunk, is_marker) — split() always yields exactly
    # one more piece than findall(), so the two zip up without a length check.
    pieces: list[tuple[str, bool]] = []
    for plain, marker in zip(between, markers):
        if plain:
            pieces.append((plain, False))
        pieces.append((marker, True))
    if between[-1]:
        pieces.append((between[-1], False))

    element = run._element
    parent = element.getparent()
    position = list(parent).index(element)

    replacements: list[tuple[object, bool]] = []
    for chunk, is_marker in pieces:
        clone = copy.deepcopy(element)
        for text_el in clone.findall(qn("w:t")):
            clone.remove(text_el)
        text_el = clone.makeelement(qn("w:t"), {qn("xml:space"): "preserve"})
        text_el.text = chunk
        clone.append(text_el)
        parent.insert(position, clone)
        position += 1
        replacements.append((clone, is_marker))
    parent.remove(element)

    highlighted = 0
    for clone, is_marker in replacements:
        if is_marker:
            Run(clone, run._parent).font.highlight_color = WD_COLOR_INDEX.YELLOW
            highlighted += 1
    return highlighted


def _highlight_review_markers(document: Document) -> int:
    """Mark every unfilled field yellow, wherever it landed.

    Headers and footers are separate stories in the XML and are walked
    explicitly — a template that puts the client name in a running header is
    normal, and a marker there would otherwise be the one the operator misses.
    """
    scopes = list(_iter_paragraph_containers(document))
    for section in document.sections:
        for part in (
            section.header,
            section.footer,
            section.first_page_header,
            section.first_page_footer,
            section.even_page_header,
            section.even_page_footer,
        ):
            scopes.extend(_iter_paragraph_containers(part))

    return sum(
        _split_run_on_markers(run)
        for scope in scopes
        for paragraph in scope.paragraphs
        # list() because _split_run_on_markers mutates the paragraph's children
        # underneath the live iterator.
        for run in list(paragraph.runs)
    )


def render_tm_agreement(values: dict[str, str | None]) -> tuple[bytes, list[str]]:
    """Fill the T&M template. Returns the .docx bytes and the labels left open.

    Never raises on a missing value. Producing a draft with terms still to
    negotiate is the normal case, not an error state — an unfilled field becomes
    a visible marker and its label comes back in the second return value so the
    caller can tell the operator what to look at. The only failure here is a
    template that is absent or malformed, and `tests/test_contract_template_tags.py`
    catches malformed at build time.

    `values` may omit tags entirely; omitted and empty mean the same thing.
    """
    if not TM_TEMPLATE.exists():
        raise TemplateMissing(
            f"Contract template not found at {TM_TEMPLATE}. It is a versioned "
            f"repo asset — see app/services/contracts/templates/README.md."
        )

    context: dict[str, str] = {}
    flagged: list[str] = []
    for field in TM_FIELDS:
        value = values.get(field.tag)
        if _blank(value):
            context[field.tag] = review_marker(field.label)
            flagged.append(field.label)
        else:
            context[field.tag] = value.strip()

    # Fresh per call. `DocxTemplate` renders by mutating its own in-memory tree,
    # so a module-level instance would hand the second caller a document with
    # the first caller's client name already baked in.
    template = _ContractTemplate(TM_TEMPLATE)
    template.render(context, autoescape=True)

    rendered = BytesIO()
    template.save(rendered)

    # Re-open rather than reaching into DocxTemplate's tree: the highlight pass
    # is a plain python-docx operation on a finished document, and keeping it
    # that way is what makes it testable on its own.
    document = Document(BytesIO(rendered.getvalue()))
    _highlight_review_markers(document)

    final = BytesIO()
    document.save(final)
    return final.getvalue(), flagged
