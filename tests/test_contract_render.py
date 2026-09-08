"""Rendering a contract template.

Tested against a template built here rather than the operator's real .docx, for
two reasons: the real one is not in the repo until they supply it, and a
synthetic one can be made to contain the *hard* cases on purpose — a tag inside
a directly-formatted run, a tag sharing a run with boilerplate, a tag in a
table, a tag in a header. Those are the cases that decide whether "matches my
formatting" is true, and they would be present in a real template only by luck.

`test_real_template_renders` picks up the operator's document once it lands.
"""
from __future__ import annotations

import re
import zipfile
from io import BytesIO

import pytest
from docx import Document
from docx.enum.text import WD_COLOR_INDEX
from docx.shared import Pt

from app.services.contracts import render as render_module
from app.services.contracts.fields import TM_FIELDS, engagement_fields, review_marker
from app.services.contracts.render import TM_TEMPLATE, render_tm_agreement

# A value that used to corrupt the document: the `&` vanished and `<Holdings>`
# swallowed the rest of the XML, table included. An entirely ordinary entity
# name, which is why it is the one used throughout.
HOSTILE_NAME = 'Smith & Co <Holdings> "Group" \'LLC\''


def _build_template(path, tags):
    """A .docx exercising every substitution position that matters."""
    doc = Document()
    tags = list(tags)

    # 1. A tag alone in a run carrying direct formatting. The fidelity case.
    para = doc.add_paragraph()
    para.add_run("Client: ")
    run = para.add_run("{{ %s }}" % tags[0])
    run.bold = True
    run.font.name = "Times New Roman"
    run.font.size = Pt(14)

    # 2. A tag sharing a run with boilerplate. The highlight-bleed case: only
    #    the marker may end up yellow, not the words around it.
    if len(tags) > 1:
        shared = doc.add_paragraph()
        shared_run = shared.add_run("State of {{ %s }}, USA" % tags[1])
        shared_run.italic = True
        shared_run.font.name = "Georgia"

    # 3. A tag in a table cell, and 4. a tag in the running header.
    if len(tags) > 2:
        doc.add_table(rows=1, cols=1).cell(0, 0).text = "{{ %s }}" % tags[2]
    if len(tags) > 3:
        doc.sections[0].header.paragraphs[0].text = "Hdr {{ %s }}" % tags[3]

    # Everything else, so the template declares the full field set.
    for tag in tags[4:]:
        doc.add_paragraph("%s: {{ %s }}" % (tag, tag))

    doc.save(path)
    return path


@pytest.fixture
def template(tmp_path, monkeypatch):
    """Point the renderer at a synthetic template covering every field."""
    path = _build_template(
        tmp_path / "tm_agreement.docx", [f.tag for f in TM_FIELDS]
    )
    monkeypatch.setattr(render_module, "TM_TEMPLATE", path)
    return path


def _all_text(content: bytes) -> str:
    """Every string in the document — body, tables, headers, footers.

    Reads the parts out of the zip rather than walking python-docx objects, so a
    stray `{{` hiding in a header or a footnote cannot be missed by a traversal
    that forgot to look there.
    """
    with zipfile.ZipFile(BytesIO(content)) as archive:
        return "".join(
            archive.read(name).decode("utf-8")
            for name in archive.namelist()
            if name.startswith("word/") and name.endswith(".xml")
        )


def _runs(content: bytes):
    doc = Document(BytesIO(content))
    for para in doc.paragraphs:
        for run in para.runs:
            yield run
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                for para in cell.paragraphs:
                    for run in para.runs:
                        yield run


def _values(**overrides):
    values = {f.tag: f"value for {f.tag}" for f in TM_FIELDS}
    values.update(overrides)
    return values


# ── The {{r … }} prefix ─────────────────────────────────────────────────────
#
# A whole section, because a real template arrives full of `{{r }}` tags —
# docxtpl's own documentation recommends the form — and the bug it caused was
# invisible in every other test here. Templates built with plain tags never
# exercise this path.

def _build_rich_text_template(path, tags):
    """A template using `{{r … }}`, the form docxtpl's docs steer you toward."""
    doc = Document()
    for i, tag in enumerate(tags):
        para = doc.add_paragraph()
        para.add_run("Label: ")
        run = para.add_run("{{r %s }}" % tag)
        # Formatting on the first tag, so the fidelity assertion has something
        # to check.
        if i == 0:
            run.bold = True
            run.font.name = "DM Sans"
    doc.save(path)
    return path


@pytest.fixture
def rich_text_template(tmp_path, monkeypatch):
    path = _build_rich_text_template(
        tmp_path / "tm_rich.docx", [f.tag for f in TM_FIELDS]
    )
    monkeypatch.setattr(render_module, "TM_TEMPLATE", path)
    return path


def test_rich_text_prefix_is_accepted(rich_text_template):
    content, flagged = render_tm_agreement(_values())
    assert flagged == []
    text = _all_text(content)
    assert "{{" not in text
    assert "value for %s" % TM_FIELDS[0].tag in text


def test_rich_text_prefix_keeps_run_formatting(rich_text_template):
    """The reason the prefix is stripped rather than honoured.

    docxtpl's `RichText` would replace the tag's run with one of its own and
    lose the bold and the font.
    """
    first = TM_FIELDS[0]
    content, _ = render_tm_agreement(_values(**{first.tag: "Acme Industries, LLC"}))

    match = [r for r in _runs(content) if r.text == "Acme Industries, LLC"]
    assert match, "the value is not a run of its own"
    assert match[0].bold is True
    assert match[0].font.name == "DM Sans"


def test_rich_text_prefix_markers_land_inside_a_text_element(rich_text_template):
    """The regression this section exists for.

    Stripping the prefix *after* docxtpl's patch pass left the value as raw text
    between two empty runs — `</w:r>[REVIEW: …]<w:r>` — which is in the file and
    invisible in Word. Two assertions, because the first one alone passed while
    the bug was live: the text was present, just not in a run.
    """
    field = TM_FIELDS[0]
    content, _ = render_tm_agreement(_values(**{field.tag: ""}))
    marker = review_marker(field.label)

    assert marker in _all_text(content)
    assert "</w:r>[REVIEW" not in _all_text(content), (
        "marker landed as raw text between runs — Word will not render it"
    )
    assert marker in [r.text for r in _runs(content)], (
        "marker is in the document but not inside any run"
    )


def test_rich_text_prefix_markers_are_highlighted(rich_text_template):
    field = TM_FIELDS[0]
    content, _ = render_tm_agreement(_values(**{field.tag: ""}))

    highlighted = [
        r.text for r in _runs(content)
        if r.font.highlight_color == WD_COLOR_INDEX.YELLOW
    ]
    assert highlighted == [review_marker(field.label)]


# ── Fully populated ─────────────────────────────────────────────────────────

def test_full_values_leave_no_tags_and_no_markers(template):
    content, flagged = render_tm_agreement(_values())

    assert flagged == []
    text = _all_text(content)
    assert "{{" not in text and "}}" not in text
    assert "[REVIEW:" not in text
    # Reopens cleanly, i.e. it is a valid .docx rather than merely bytes.
    assert Document(BytesIO(content)).paragraphs


def test_direct_run_formatting_survives(template):
    """The reason this feature renders a template instead of building a doc."""
    first = TM_FIELDS[0]
    content, _ = render_tm_agreement(_values(**{first.tag: "Acme Industries, LLC"}))

    match = [r for r in _runs(content) if r.text == "Acme Industries, LLC"]
    assert match, "the substituted value is not a run of its own"
    run = match[0]
    assert run.bold is True
    assert run.font.name == "Times New Roman"
    assert run.font.size == Pt(14)


def test_xml_hostile_value_does_not_corrupt_the_document(template):
    """`&` and `<` in a value used to swallow the document's structure."""
    first = TM_FIELDS[0]
    content, _ = render_tm_agreement(_values(**{first.tag: HOSTILE_NAME}))

    doc = Document(BytesIO(content))
    texts = [r.text for r in _runs(content)]
    assert HOSTILE_NAME in texts, f"value did not round-trip; got {texts}"
    # The table lives after the hostile value in document order, so it is the
    # canary for structural damage.
    assert doc.tables, "the table after the substitution was destroyed"


def test_renders_into_tables_and_headers(template):
    content, _ = render_tm_agreement(_values())
    doc = Document(BytesIO(content))

    cell_text = doc.tables[0].cell(0, 0).text
    assert "{{" not in cell_text and cell_text.strip()

    header_text = doc.sections[0].header.paragraphs[0].text
    assert "{{" not in header_text and header_text.strip() != "Hdr"


# ── Blanks ──────────────────────────────────────────────────────────────────

@pytest.mark.parametrize("blank", ["", "   ", None])
def test_blank_becomes_a_highlighted_marker(template, blank):
    """Empty, whitespace, and absent all mean the same thing."""
    field = TM_FIELDS[0]
    content, flagged = render_tm_agreement(_values(**{field.tag: blank}))

    assert flagged == [field.label]
    marker = review_marker(field.label)
    assert marker in _all_text(content)

    highlighted = [
        r for r in _runs(content)
        if r.font.highlight_color == WD_COLOR_INDEX.YELLOW
    ]
    assert [r.text for r in highlighted] == [marker]


def test_omitted_key_is_treated_as_blank(template):
    """`values` need not mention every tag."""
    values = _values()
    field = TM_FIELDS[0]
    del values[field.tag]

    content, flagged = render_tm_agreement(values)
    assert flagged == [field.label]
    assert review_marker(field.label) in _all_text(content)


def test_marker_keeps_surrounding_formatting_and_does_not_bleed(template):
    """A tag sharing a run with boilerplate must highlight only the marker.

    Highlighting the whole run would paint "State of" and ", USA" yellow, which
    reads as though the boilerplate needs review rather than the value.
    """
    field = TM_FIELDS[1]
    content, _ = render_tm_agreement(_values(**{field.tag: ""}))
    marker = review_marker(field.label)

    by_text = {r.text: r for r in _runs(content)}
    assert marker in by_text, f"marker run not found among {list(by_text)}"

    assert by_text[marker].font.highlight_color == WD_COLOR_INDEX.YELLOW
    # rPr came across with the split, so the marker still reads as body text.
    assert by_text[marker].italic is True
    assert by_text[marker].font.name == "Georgia"

    for neighbour in ("State of ", ", USA"):
        assert by_text[neighbour].font.highlight_color is None
        assert by_text[neighbour].italic is True


def test_every_blank_field_is_flagged_and_marked(template):
    """All of them at once — the "generate an empty draft to negotiate" case."""
    content, flagged = render_tm_agreement({})

    assert flagged == [f.label for f in TM_FIELDS]
    text = _all_text(content)
    for field in TM_FIELDS:
        assert review_marker(field.label) in text
    assert "{{" not in text


def test_marker_in_a_header_is_highlighted(template):
    """Headers are a separate story in the XML and a separate walk in the code."""
    field = TM_FIELDS[3]
    content, _ = render_tm_agreement(_values(**{field.tag: ""}))

    with zipfile.ZipFile(BytesIO(content)) as archive:
        headers = [
            archive.read(n).decode()
            for n in archive.namelist()
            if re.match(r"word/header\d*\.xml", n)
        ]
    assert any(
        review_marker(field.label) in h and "w:highlight" in h for h in headers
    ), "the header marker was rendered but not highlighted"


# ── Isolation and failure ───────────────────────────────────────────────────

def test_consecutive_renders_do_not_leak(template):
    """`DocxTemplate` renders by mutating its own tree.

    A cached module-level instance would hand the second caller a document with
    the first caller's client name already substituted in.
    """
    field = TM_FIELDS[0]
    first, _ = render_tm_agreement(_values(**{field.tag: "First Client, LLC"}))
    second, _ = render_tm_agreement(_values(**{field.tag: "Second Client, LLC"}))

    assert "First Client, LLC" not in _all_text(second)
    assert "Second Client, LLC" in _all_text(second)
    assert "First Client, LLC" in _all_text(first)


def test_missing_template_raises_template_missing(tmp_path, monkeypatch):
    monkeypatch.setattr(render_module, "TM_TEMPLATE", tmp_path / "absent.docx")
    with pytest.raises(render_module.TemplateMissing, match="README"):
        render_tm_agreement(_values())


def test_engagement_fields_are_declared():
    """A T&M contract has to vary by *something* per deal."""
    assert engagement_fields(), (
        "TM_FIELDS declares no engagement fields, so every draft for a given "
        "client would be byte-identical"
    )


# ── The operator's real template, once it exists ────────────────────────────

def test_real_template_renders():
    if not TM_TEMPLATE.exists():
        pytest.skip(
            f"No contract template at {TM_TEMPLATE} yet — see "
            f"app/services/contracts/templates/README.md."
        )

    content, flagged = render_tm_agreement(_values())
    assert flagged == []
    text = _all_text(content)
    assert "{{" not in text, "an unfilled tag survived into the rendered draft"
    assert "[REVIEW:" not in text
